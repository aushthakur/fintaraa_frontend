import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import csv
import os

BASE_URL = "https://fintaraa-frontend-2o45.onrender.com"

data = [
    # ── 1. PERSONAL LOANS ──────────────────────────────────────────────────────────
    {
        "category": "Personal Loans",
        "page_name": "Personal Loan",
        "tagline": "Unsecured personal credit for immediate financial needs with instant paperless approval.",
        "path": "/products/personal-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Personal Loans",
        "page_name": "Instant Personal Loan",
        "tagline": "Fast digital application with 10-minute automated algorithmic sanction and disbursal.",
        "path": "/products/instant-loan",
        "type": "Marketplace Page",
        "status": "Active"
    },
    {
        "category": "Personal Loans",
        "page_name": "Education Loan",
        "tagline": "Study in India or overseas with 100% tuition coverage, living expenses & zero collateral up to ₹50L.",
        "path": "/products/education-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Personal Loans",
        "page_name": "Renovation Loan",
        "tagline": "Home upgrade financing for modular kitchens, architectural remodeling, and interior design.",
        "path": "/products/renovation-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Personal Loans",
        "page_name": "Wedding Loan",
        "tagline": "Marriage celebration expenses covering destination venue booking, jewelry, catering & bridal decor.",
        "path": "/products/wedding-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Personal Loans",
        "page_name": "Marriage Loan (Alias)",
        "tagline": "Dedicated wedding celebration and event expense financing.",
        "path": "/products/marriage-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Personal Loans",
        "page_name": "Medical Loan",
        "tagline": "Healthcare emergency support for surgeries, hospital treatments, dental, IVF & elective procedures.",
        "path": "/products/medical-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Personal Loans",
        "page_name": "Healthcare Loan (Alias)",
        "tagline": "Medical treatments and specialized hospital procedure financing.",
        "path": "/products/healthcare-loan",
        "type": "Product Page",
        "status": "Active"
    },

    # ── 2. HOME & PROPERTY LOANS ──────────────────────────────────────────────────
    {
        "category": "Home & Property",
        "page_name": "Home Loan",
        "tagline": "Housing purchase loan for new apartments, builder floors, ready-to-move-in flats & resale properties.",
        "path": "/products/home-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Home & Property",
        "page_name": "Housing Loan (Alias)",
        "tagline": "Competitive residential mortgage financing.",
        "path": "/products/housing-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Home & Property",
        "page_name": "Balance Transfer & Top-Up Loan",
        "tagline": "Lower your EMI by transferring home loan at low rates from 8.35% p.a. + high-value top-up funds.",
        "path": "/products/balance-transfer-top-up-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Home & Property",
        "page_name": "Balance Transfer Loan",
        "tagline": "Reduce existing loan burden and switch to prime lenders.",
        "path": "/products/balance-transfer-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Home & Property",
        "page_name": "Top-Up Loan",
        "tagline": "Instant cash credit added directly on existing mortgage sanction.",
        "path": "/products/top-up-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Home & Property",
        "page_name": "Loan Against Property (LAP)",
        "tagline": "Unlock liquidity from residential or commercial property up to ₹25 Cr with tenures up to 20 years.",
        "path": "/products/loan-against-property",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Home & Property",
        "page_name": "LAP (Alias)",
        "tagline": "Property equity monetization and overdraft credit line.",
        "path": "/products/lap",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Home & Property",
        "page_name": "Construction Loan",
        "tagline": "Residential self-construction, plot + building composite packages, and floor additions.",
        "path": "/products/construction-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Home & Property",
        "page_name": "Commercial Purchases Loan",
        "tagline": "Commercial property needs covering office spaces, industrial premises, and retail shops.",
        "path": "/products/commercial-purchases-loan",
        "type": "Product Page",
        "status": "Active"
    },

    # ── 3. BUSINESS LOANS ─────────────────────────────────────────────────────────
    {
        "category": "Business Loans",
        "page_name": "Business Loan",
        "tagline": "Commercial growth capital up to ₹75 Lakhs collateral-free with fast 48-hour disbursal.",
        "path": "/products/business-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Business Loans",
        "page_name": "Working Capital Loan",
        "tagline": "Day-to-day cash flow management, vendor invoice clearance, and inventory seasonal spikes.",
        "path": "/products/working-capital-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Business Loans",
        "page_name": "OD / Overdraft Loan",
        "tagline": "Flexible credit limit attached to current account where interest is charged only on utilized funds.",
        "path": "/products/od-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Business Loans",
        "page_name": "DOD Loan (Drop-line Overdraft)",
        "tagline": "Structured drop-line overdraft facility with monthly reducing principal limits.",
        "path": "/products/dod-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Business Loans",
        "page_name": "Machinery & Equipment Loan",
        "tagline": "Equipment & machinery financing for manufacturing CNC units, automation systems, and heavy tools.",
        "path": "/products/machinery-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Business Loans",
        "page_name": "Industrial Loan",
        "tagline": "Large-scale industrial expansion, warehouse construction, and manufacturing plant credit.",
        "path": "/products/industrial-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Business Loans",
        "page_name": "Agriculture & Farm Loan",
        "tagline": "Farming infrastructure, tractor financing, solar pump installation, and crop cultivation credit.",
        "path": "/products/agriculture-loan",
        "type": "Product Page",
        "status": "Active"
    },

    # ── 4. VEHICLE LOANS ──────────────────────────────────────────────────────────
    {
        "category": "Vehicle Loans",
        "page_name": "Car Loan",
        "tagline": "New car financing with up to 100% on-road funding and concessional interest rates for electric vehicles.",
        "path": "/products/car-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Vehicle Loans",
        "page_name": "Used Car Loan",
        "tagline": "Pre-owned car financing with free RC transfer and valuation assistance up to 150% valuation.",
        "path": "/products/used-car-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Vehicle Loans",
        "page_name": "Two-Wheeler Loan",
        "tagline": "Two-wheeler financing for daily commuter motorcycles, sports bikes, and smart EV scooters.",
        "path": "/products/two-wheeler-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Vehicle Loans",
        "page_name": "Loan Against Car",
        "tagline": "Liquidity against car equity without surrendering physical vehicle usage.",
        "path": "/products/loan-against-car",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Vehicle Loans",
        "page_name": "Loan Against Car Value",
        "tagline": "Specialized vehicle refinance and equity top-up monetization.",
        "path": "/products/loan-against-car-value",
        "type": "Product Page",
        "status": "Active"
    },

    # ── 5. LOANS AGAINST ASSETS ───────────────────────────────────────────────────
    {
        "category": "Loans Against Assets",
        "page_name": "Gold Loan",
        "tagline": "Instant doorstep or branch liquidity pledged against 18K–24K gold jewelry at lowest 8.40% p.a. rate.",
        "path": "/products/gold-loan",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Loans Against Assets",
        "page_name": "Loan Against Securities (LAS)",
        "tagline": "Instant credit line pledged against Demat mutual funds, stocks, and bonds without selling assets.",
        "path": "/products/loan-against-security",
        "type": "Product Page",
        "status": "Active"
    },
    {
        "category": "Loans Against Assets",
        "page_name": "Solar Rooftop Loan",
        "tagline": "Government subsidized rooftop solar installation loans under PM Surya Ghar Muft Bijli Yojana.",
        "path": "/products/solar-loan",
        "type": "Product Page",
        "status": "Active"
    },

    # ── 6. CREDIT CARDS ───────────────────────────────────────────────────────────
    {
        "category": "Credit Cards",
        "page_name": "Credit Cards Marketplace",
        "tagline": "Compare 50+ credit cards across lifetime-free, cashback, travel, and premium reward programs.",
        "path": "/credit-cards",
        "type": "Marketplace Page",
        "status": "Active"
    },
    {
        "category": "Credit Cards",
        "page_name": "Credit Card (Alias)",
        "tagline": "Credit cards directory and instant eligibility check.",
        "path": "/credit-card",
        "type": "Marketplace Page",
        "status": "Active"
    },
    {
        "category": "Credit Cards",
        "page_name": "Cashback Credit Cards",
        "tagline": "Best credit cards for flat cashbacks on groceries, utility bills, Swiggy, and Amazon shopping.",
        "path": "/credit-cards/cashback",
        "type": "Category Hub",
        "status": "Active"
    },
    {
        "category": "Credit Cards",
        "page_name": "Fuel Credit Cards",
        "tagline": "Cards offering 1% fuel surcharge waivers and bonus reward points at Indian Oil, HPCL, and BPCL.",
        "path": "/credit-cards/fuel",
        "type": "Category Hub",
        "status": "Active"
    },
    {
        "category": "Credit Cards",
        "page_name": "Travel Credit Cards",
        "tagline": "Free airport lounge access, international concierge, low forex markup & complimentary airmiles.",
        "path": "/credit-cards/travel",
        "type": "Category Hub",
        "status": "Active"
    },
    {
        "category": "Credit Cards",
        "page_name": "Rewards Credit Cards",
        "tagline": "Accelerated reward points cards for lifestyle, weekend fine dining, luxury retail, and gift vouchers.",
        "path": "/credit-cards/rewards",
        "type": "Category Hub",
        "status": "Active"
    },

    # ── 7. INSURANCE SOLUTIONS ────────────────────────────────────────────────────
    {
        "category": "Insurance",
        "page_name": "Health Insurance",
        "tagline": "Comprehensive cashless medical cover, critical illness riders, and family floater policies.",
        "path": "/products/health-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Term Life Insurance",
        "tagline": "High-cover affordable term life insurance with whole-life benefits and return-of-premium options.",
        "path": "/products/term-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Life Insurance",
        "tagline": "Endowment, money-back, and wealth creation life insurance policies with tax savings under 80C.",
        "path": "/products/life-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Car & Motor Insurance",
        "tagline": "Zero-depreciation comprehensive four-wheeler motor insurance with instant cashless claims.",
        "path": "/products/car-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Two Wheeler Bike Insurance",
        "tagline": "Instant online bike insurance renewal with third-party and comprehensive accident covers.",
        "path": "/products/bike-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Commercial Vehicle Insurance",
        "tagline": "Fleet and commercial truck, bus, and taxi insurance policies.",
        "path": "/products/vehicle-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Home & Property Insurance",
        "tagline": "Protection against fire, burglary, natural calamities, and structural damage for homes.",
        "path": "/products/home-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Commercial Property Insurance",
        "tagline": "Business premises, factory, and commercial warehouse asset protection.",
        "path": "/products/property-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Shop & Retail Insurance",
        "tagline": "Retail shop package cover protecting cash in transit, inventory stock, and burglary.",
        "path": "/products/shop-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Stock & Inventory Insurance",
        "tagline": "Coverage against warehouse fire, water damage, spoilage, and transit theft.",
        "path": "/products/stock-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Machinery Breakdown Insurance",
        "tagline": "Protection against mechanical breakdown, electrical failure, and industrial equipment repair costs.",
        "path": "/products/machinery-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Critical Illness Insurance",
        "tagline": "Lump sum financial benefit on diagnosis of 36+ major critical illnesses.",
        "path": "/products/critical-illness-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Personal Accident Insurance",
        "tagline": "Worldwide cover for accidental death, permanent total disability, and income loss.",
        "path": "/products/personal-accident-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Travel Insurance",
        "tagline": "International and domestic travel protection for medical emergencies, trip cancellations, and lost baggage.",
        "path": "/products/travel-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Retirement & Pension Plan",
        "tagline": "Guaranteed annuity pension plans and wealth preservation retirement solutions.",
        "path": "/products/retirement-plan",
        "type": "Insurance Page",
        "status": "Active"
    },
    {
        "category": "Insurance",
        "page_name": "Group Health & Employee Insurance",
        "tagline": "Corporate employee group medical insurance and accidental health schemes.",
        "path": "/products/group-insurance",
        "type": "Insurance Page",
        "status": "Active"
    },

    # ── 8. CALCULATORS & FINANCIAL TOOLS ──────────────────────────────────────────
    {
        "category": "Calculators & Tools",
        "page_name": "Calculators Hub",
        "tagline": "Financial planning tools suite for loan EMIs, interest schedules, and repayment models.",
        "path": "/calculators",
        "type": "Calculator Hub",
        "status": "Active"
    },
    {
        "category": "Calculators & Tools",
        "page_name": "Loan EMI Calculator",
        "tagline": "Real-time monthly installment calculator with amortization breakdown and interest graphs.",
        "path": "/calculators/emi-calculator",
        "type": "Tool Page",
        "status": "Active"
    },
    {
        "category": "Calculators & Tools",
        "page_name": "Home Loan EMI Calculator",
        "tagline": "Calculate home loan installments, tenure savings, and tax deduction under 80C & 24b.",
        "path": "/calculators/home-loan-calculator",
        "type": "Tool Page",
        "status": "Active"
    },
    {
        "category": "Calculators & Tools",
        "page_name": "Personal Loan Calculator",
        "tagline": "Compute personal loan installments and repayment timelines.",
        "path": "/calculators/personal-loan-calculator",
        "type": "Tool Page",
        "status": "Active"
    },
    {
        "category": "Calculators & Tools",
        "page_name": "Loan Eligibility Calculator",
        "tagline": "Estimate maximum borrowing capacity based on net monthly salary and existing obligations.",
        "path": "/calculators/loan-eligibility-calculator",
        "type": "Tool Page",
        "status": "Active"
    },
    {
        "category": "Calculators & Tools",
        "page_name": "Financial Tools Hub",
        "tagline": "Suite of business valuation, tax, and loan eligibility utilities.",
        "path": "/tools",
        "type": "Tool Hub",
        "status": "Active"
    },
    {
        "category": "Calculators & Tools",
        "page_name": "Free CIBIL Score Checker",
        "tagline": "Instant free credit score check powered by official credit bureau API without hitting score.",
        "path": "/cibil-score",
        "type": "Utility Page",
        "status": "Active"
    },
    {
        "category": "Calculators & Tools",
        "page_name": "Comprehensive CIBIL Report",
        "tagline": "Full credit score report analysis, credit factors, dispute assistance, and improvement guide.",
        "path": "/cibil-score/report",
        "type": "Utility Page",
        "status": "Active"
    },

    # ── 9. TAX & CORPORATE COMPLIANCE ─────────────────────────────────────────────
    {
        "category": "Corporate & Tax Services",
        "page_name": "Company Registration",
        "tagline": "Private Limited, LLP, OPC, and Partnership firm incorporation with MCA & RoC filings.",
        "path": "/company-registration",
        "type": "Service Page",
        "status": "Active"
    },
    {
        "category": "Corporate & Tax Services",
        "page_name": "GST Registration",
        "tagline": "New GSTIN registration, modification, composition scheme, and monthly GSTR filing support.",
        "path": "/gst-registration",
        "type": "Service Page",
        "status": "Active"
    },
    {
        "category": "Corporate & Tax Services",
        "page_name": "ITR Filing Services",
        "tagline": "Expert chartered accountant assisted income tax return e-filing for salaried and business owners.",
        "path": "/itr-filing",
        "type": "Service Page",
        "status": "Active"
    },
    {
        "category": "Corporate & Tax Services",
        "page_name": "MSME / Udyam Registration",
        "tagline": "Government MSME Udyam certification for priority bank lending and subsidy schemes.",
        "path": "/msme-registration",
        "type": "Service Page",
        "status": "Active"
    },
    {
        "category": "Corporate & Tax Services",
        "page_name": "ROC Filing & Compliance",
        "tagline": "Annual MCA returns, AOC-4, MGT-7, director KYC, and statutory filings.",
        "path": "/roc-filing",
        "type": "Service Page",
        "status": "Active"
    },
    {
        "category": "Corporate & Tax Services",
        "page_name": "Annual Compliance Package",
        "tagline": "End-to-end secretarial, accounting, tax audit, and statutory compliance bundle.",
        "path": "/annual-compliance",
        "type": "Service Page",
        "status": "Active"
    },
    {
        "category": "Corporate & Tax Services",
        "page_name": "CMA Data & Project Report",
        "tagline": "Bank-ready detailed project reports (DPR), CMA data, and financial projections for bank loans.",
        "path": "/project-report",
        "type": "Service Page",
        "status": "Active"
    },
    {
        "category": "Corporate & Tax Services",
        "page_name": "Tax Compliance & Advisory",
        "tagline": "Direct and indirect corporate tax optimization, notice assistance, and TDS compliance.",
        "path": "/tax-compliance",
        "type": "Service Page",
        "status": "Active"
    },

    # ── 10. PARTNERSHIP & NETWORK ─────────────────────────────────────────────────
    {
        "category": "Partnership & Growth",
        "page_name": "Become a DSA Partner",
        "tagline": "Join India's highest paying Direct Selling Agent network with instant commissions and CRM tools.",
        "path": "/become-dsa",
        "type": "Partnership Page",
        "status": "Active"
    },
    {
        "category": "Partnership & Growth",
        "page_name": "Franchise Program",
        "tagline": "Start your own Fintaraa financial services franchise office with marketing and branch support.",
        "path": "/franchise",
        "type": "Partnership Page",
        "status": "Active"
    },
    {
        "category": "Partnership & Growth",
        "page_name": "Partner Portal",
        "tagline": "Channel partner registration, lead tracking, payout calculation, and training modules.",
        "path": "/partner",
        "type": "Portal Page",
        "status": "Active"
    },
    {
        "category": "Partnership & Growth",
        "page_name": "Partner Login",
        "tagline": "Secure CRM dashboard access for registered Fintaraa DSA channel partners.",
        "path": "/partner/login",
        "type": "Portal Page",
        "status": "Active"
    },
    {
        "category": "Partnership & Growth",
        "page_name": "Refer and Earn Program",
        "tagline": "Refer friends and colleagues for loans or credit cards and earn cash rewards.",
        "path": "/refer-and-earn",
        "type": "Reward Program",
        "status": "Active"
    },
    {
        "category": "Partnership & Growth",
        "page_name": "All Lending Partners",
        "tagline": "Complete directory of 40+ empaneled banking and NBFC institutional partners.",
        "path": "/partners/all",
        "type": "Directory Page",
        "status": "Active"
    },
    {
        "category": "Partnership & Growth",
        "page_name": "Partners By Product",
        "tagline": "Interactive partner directory categorized by loan, credit card, and insurance product verticals.",
        "path": "/partners-by-product",
        "type": "Directory Page",
        "status": "Active"
    },
    {
        "category": "Partnership & Growth",
        "page_name": "Bank Partners Overview",
        "tagline": "Comprehensive directory of bank profiles, interest rate cards, and city branches.",
        "path": "/banks",
        "type": "Directory Page",
        "status": "Active"
    },

    # ── 11. COMPANY, KNOWLEDGE & SUPPORT ──────────────────────────────────────────
    {
        "category": "Company & Resources",
        "page_name": "About Fintaraa",
        "tagline": "Learn about Fintaraa's mission, leadership team, vision, and fintech growth story.",
        "path": "/about-us",
        "type": "Corporate Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Contact Us",
        "tagline": "Customer care helpline, email support, corporate office addresses, and grievance desk.",
        "path": "/contact-us",
        "type": "Support Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Customer Support Hub",
        "tagline": "Self-service support desk, FAQs, complaint tickets, and loan officer callback requests.",
        "path": "/support",
        "type": "Support Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Track Application Status",
        "tagline": "Real-time loan and credit card application tracker using mobile OTP or reference ID.",
        "path": "/application-status",
        "type": "Utility Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Knowledge Hub",
        "tagline": "Financial education guides, borrowing handbooks, RBI regulatory updates, and glossary.",
        "path": "/knowledge-hub",
        "type": "Content Hub",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Financial Blog",
        "tagline": "Articles, market analysis, interest rate forecasts, and personal finance tips.",
        "path": "/blog",
        "type": "Content Hub",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "All Blog Articles",
        "tagline": "Comprehensive library of all published financial and loan advisory blogs.",
        "path": "/blog/all",
        "type": "Content Hub",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Exclusive Offers & Deals",
        "tagline": "Curated zero processing fee offers, cashback bonuses, and special interest discounts.",
        "path": "/offers",
        "type": "Deals Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Careers at Fintaraa",
        "tagline": "Explore open job positions across engineering, credit underwriting, sales, and operations.",
        "path": "/careers",
        "type": "Corporate Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Frequently Asked Questions (FAQs)",
        "tagline": "Answers to common borrower queries on documentation, CIBIL, interest rates, and disbursals.",
        "path": "/faqs",
        "type": "Support Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Customer Feedback & Reviews",
        "tagline": "Share your borrowing experience and suggest platform improvements.",
        "path": "/feedback",
        "type": "Support Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Customer Testimonials",
        "tagline": "Verified stories and reviews from 100,000+ satisfied borrowers across India.",
        "path": "/testimonials",
        "type": "Trust Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Video Testimonials",
        "tagline": "Watch customer video success stories and entrepreneur loan case studies.",
        "path": "/video-testimonials",
        "type": "Trust Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Awards and Recognitions",
        "tagline": "Industry fintech awards, excellence recognitions, and media honors awarded to Fintaraa.",
        "path": "/awards-and-recognitions",
        "type": "Trust Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Press & Media Coverage",
        "tagline": "Fintaraa in top financial news publications, media coverage, and leader interviews.",
        "path": "/press",
        "type": "Media Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Press Releases",
        "tagline": "Official company press announcements, funding milestones, and partnership releases.",
        "path": "/press-release",
        "type": "Media Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Privacy Policy",
        "tagline": "Data security protocols, encryption standards, and customer information privacy commitment.",
        "path": "/privacy-policy",
        "type": "Legal Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Terms and Conditions",
        "tagline": "Terms of platform use, loan marketplace terms, and user agreements.",
        "path": "/terms-and-conditions",
        "type": "Legal Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Grievance Redressal",
        "tagline": "RBI compliant grievance redressal cell and designated grievance officer contact details.",
        "path": "/grievance",
        "type": "Legal Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Loan Disclosure & Fair Practices",
        "tagline": "Statutory APR rate band disclosures, repayment tenures, and partner lending charters.",
        "path": "/loan-disclosure",
        "type": "Legal Page",
        "status": "Active"
    },
    {
        "category": "Company & Resources",
        "page_name": "Delete Account & Data Rights",
        "tagline": "User data deletion and GDPR/DPDP compliant right to be forgotten request portal.",
        "path": "/delete-account",
        "type": "Legal Page",
        "status": "Active"
    },
]

# ── Create OpenPyXL Workbook ──────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Fintaraa Major Pages Directory"

# Styling definitions
title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
subtitle_font = Font(name="Calibri", size=10, italic=True, color="E2E8F0")
title_fill = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid") # Deep Royal Purple

header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="6424C7", end_color="6424C7", fill_type="solid") # Fintaraa Signature Purple

category_font = Font(name="Calibri", size=11, bold=True, color="4C1D95")
category_fill = PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid") # Light Purple Pill

row_font = Font(name="Calibri", size=10, color="0F172A")
link_font = Font(name="Calibri", size=10, color="6424C7", underline="single")
status_active_font = Font(name="Calibri", size=10, bold=True, color="059669") # Emerald Green
zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

# 1. Title Banner
ws.merge_cells("A1:G1")
title_cell = ws["A1"]
title_cell.value = "Fintaraa - Master Pages & URLs Directory"
title_cell.font = title_font
title_cell.fill = title_fill
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 36

# Subtitle Banner
ws.merge_cells("A2:G2")
sub_cell = ws["A2"]
sub_cell.value = f"Base Domain: {BASE_URL}  |  Total Pages: {len(data)}  |  Design Theme: White & Purple (#6424C7)"
sub_cell.font = subtitle_font
sub_cell.fill = title_fill
sub_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[2].height = 20

# 2. Table Headers
headers = [
    "S.No",
    "Category",
    "Page / Product Name",
    "Live Production URL",
    "Relative Route Path",
    "Page Type",
    "Tagline & Description"
]

ws.row_dimensions[4].height = 28
for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col_idx)
    cell.value = h
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center" if col_idx in [1, 6] else "left", vertical="center")
    cell.border = thin_border

# 3. Populate Rows
current_row = 5
current_cat = None

for idx, item in enumerate(data, 1):
    live_url = f"{BASE_URL}{item['path']}"
    
    # Optional Category Header Separator
    if item['category'] != current_cat:
        current_cat = item['category']
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        cat_cell = ws.cell(row=current_row, column=1)
        cat_cell.value = f"❖ {current_cat.upper()}"
        cat_cell.font = category_font
        cat_cell.fill = category_fill
        cat_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[current_row].height = 24
        current_row += 1

    ws.row_dimensions[current_row].height = 22
    is_zebra = (idx % 2 == 0)
    row_fill = zebra_fill if is_zebra else None

    # Col 1: S.No
    c1 = ws.cell(row=current_row, column=1, value=idx)
    c1.font = row_font
    c1.alignment = Alignment(horizontal="center", vertical="center")
    
    # Col 2: Category
    c2 = ws.cell(row=current_row, column=2, value=item['category'])
    c2.font = row_font
    c2.alignment = Alignment(horizontal="left", vertical="center")
    
    # Col 3: Page Name
    c3 = ws.cell(row=current_row, column=3, value=item['page_name'])
    c3.font = Font(name="Calibri", size=10, bold=True, color="0F172A")
    c3.alignment = Alignment(horizontal="left", vertical="center")

    # Col 4: Live URL (Clickable Hyperlink)
    c4 = ws.cell(row=current_row, column=4, value=live_url)
    c4.hyperlink = live_url
    c4.font = link_font
    c4.alignment = Alignment(horizontal="left", vertical="center")

    # Col 5: Relative Path
    c5 = ws.cell(row=current_row, column=5, value=item['path'])
    c5.font = row_font
    c5.alignment = Alignment(horizontal="left", vertical="center")

    # Col 6: Page Type
    c6 = ws.cell(row=current_row, column=6, value=item['type'])
    c6.font = row_font
    c6.alignment = Alignment(horizontal="center", vertical="center")

    # Col 7: Tagline & Description
    c7 = ws.cell(row=current_row, column=7, value=item['tagline'])
    c7.font = row_font
    c7.alignment = Alignment(horizontal="left", vertical="center")

    for c in [c1, c2, c3, c4, c5, c6, c7]:
        c.border = thin_border
        if row_fill:
            c.fill = row_fill

    current_row += 1

# 4. Set Column Widths
col_widths = {
    "A": 8,   # S.No
    "B": 24,  # Category
    "C": 32,  # Page Name
    "D": 65,  # Live URL
    "E": 36,  # Relative Route
    "F": 20,  # Page Type
    "G": 75   # Tagline / Description
}

for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

# Enable Filters
ws.auto_filter.ref = f"A4:G{current_row-1}"

# Save Excel Files
excel_paths = [
    "/Users/apple/Downloads/Fintaraa_backup/Fintaraa_backup/fintaraa-website/Fintaraa_Major_Pages_URL_Directory.xlsx",
    "/Users/apple/Downloads/Fintaraa_Major_Pages_URL_Directory.xlsx"
]

for p in excel_paths:
    os.makedirs(os.path.dirname(p), exist_ok=True)
    wb.save(p)
    print(f"Saved Excel file to: {p}")

# Also generate corresponding CSV file
csv_paths = [
    "/Users/apple/Downloads/Fintaraa_backup/Fintaraa_backup/fintaraa-website/Fintaraa_Major_Pages_URL_Directory.csv",
    "/Users/apple/Downloads/Fintaraa_Major_Pages_URL_Directory.csv"
]

csv_rows = [["S.No", "Category", "Page / Product Name", "Live Production URL", "Relative Route Path", "Page Type", "Tagline & Description"]]
for idx, item in enumerate(data, 1):
    csv_rows.append([
        str(idx),
        item['category'],
        item['page_name'],
        f"{BASE_URL}{item['path']}",
        item['path'],
        item['type'],
        item['tagline']
    ])

for cp in csv_paths:
    os.makedirs(os.path.dirname(cp), exist_ok=True)
    with open(cp, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerows(csv_rows)
    print(f"Saved CSV file to: {cp}")
